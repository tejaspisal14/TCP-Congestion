#!/usr/bin/env python3
"""Build the revised two-column research-manuscript draft and figures."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "docs"
DOCS.mkdir(exist_ok=True)


def draw_box(ax, xy, w, h, text, fontsize=10, facecolor="#f4f6f8", edgecolor="#263746", weight="normal"):
    patch = FancyBboxPatch(xy, w, h, boxstyle="round,pad=0.02,rounding_size=0.04",
                           linewidth=1.2, edgecolor=edgecolor, facecolor=facecolor)
    ax.add_patch(patch)
    ax.text(xy[0] + w/2, xy[1] + h/2, text, ha="center", va="center", fontsize=fontsize,
            weight=weight, wrap=True, color="#17212b")
    return patch


def arrow(ax, start, end, text=None, color="#34495e", rad=0.0):
    arr = FancyArrowPatch(start, end, arrowstyle="-|>", mutation_scale=13, lw=1.2,
                          color=color, connectionstyle=f"arc3,rad={rad}")
    ax.add_patch(arr)
    if text:
        mx, my = (start[0]+end[0])/2, (start[1]+end[1])/2
        ax.text(mx, my + 0.06, text, ha="center", va="bottom", fontsize=8, color=color)


def create_figures():
    # Architecture figure
    fig, ax = plt.subplots(figsize=(10, 4.7))
    ax.set_xlim(0, 10); ax.set_ylim(0, 5); ax.axis("off")
    draw_box(ax, (3.5, 4.05), 3.0, 0.65, "Experiment controller\nScenario config • run order • repetitions", 10, "#e8eef5", weight="bold")
    draw_box(ax, (0.35, 2.45), 2.55, 0.95, "Mininet dumbbell testbed\nLinux TCP • OVS • TCLink", 9.3, "#eaf4ee", weight="bold")
    draw_box(ax, (3.72, 2.45), 2.55, 0.95, "Traffic and sampling\niperf3 • ss • raw JSON", 9.3, "#fff3df", weight="bold")
    draw_box(ax, (7.08, 2.45), 2.55, 0.95, "Persistent artifacts\nruns.csv • environment.json", 9.0, "#f5ebf5", weight="bold")
    draw_box(ax, (2.1, 0.65), 2.7, 0.9, "Analysis pipeline\nPandas • confidence intervals", 9.3, "#e8eef5", weight="bold")
    draw_box(ax, (5.45, 0.65), 2.7, 0.9, "Paper outputs\nsummary tables • plots • report", 9.3, "#eaf4ee", weight="bold")
    arrow(ax, (4.35, 4.05), (1.8, 3.4), "configure")
    arrow(ax, (2.9, 2.9), (3.72, 2.9), "traffic")
    arrow(ax, (6.27, 2.9), (7.08, 2.9), "measurements")
    arrow(ax, (8.0, 2.45), (4.2, 1.55), "CSV / raw logs", rad=-0.12)
    arrow(ax, (4.8, 1.1), (5.45, 1.1), "summarize")
    ax.set_title("Reproducible experiment and analysis pipeline", fontsize=13, weight="bold", pad=12)
    fig.tight_layout()
    fig.savefig(DOCS / "architecture.png", dpi=240, bbox_inches="tight")
    plt.close(fig)

    # Dumbbell topology figure
    fig, ax = plt.subplots(figsize=(10, 4.6))
    ax.set_xlim(0, 10); ax.set_ylim(0, 5); ax.axis("off")
    for i, y in enumerate((3.8, 2.75, 1.7), 1):
        draw_box(ax, (0.35, y), 1.35, 0.55, f"h{i}\nSender", 9, "#e8eef5")
        draw_box(ax, (8.3, y), 1.35, 0.55, f"h{i+3}\nReceiver", 9, "#eaf4ee")
        arrow(ax, (1.7, y+0.275), (2.7, 2.75), None, rad=(y-2.75)*0.03)
        arrow(ax, (7.3, 2.75), (8.3, y+0.275), None, rad=(2.75-y)*0.03)
    draw_box(ax, (2.7, 2.25), 1.25, 1.0, "s1\nOVS", 11, "#fff3df", weight="bold")
    draw_box(ax, (6.05, 2.25), 1.25, 1.0, "s2\nOVS", 11, "#fff3df", weight="bold")
    draw_box(ax, (4.18, 2.1), 1.65, 1.3, "Shared\nbottleneck\nTCLink", 10, "#f5ebf5", weight="bold")
    arrow(ax, (3.95, 2.75), (4.18, 2.75), "")
    arrow(ax, (5.83, 2.75), (6.05, 2.75), "")
    ax.text(5.0, 1.45, "Configured bandwidth • delay • random loss • queue limit", ha="center", va="center", fontsize=9)
    ax.text(5.0, 0.75, "Single-flow tests: h1 → h4 | Mixed-flow test: h1/h2/h3 → h4/h5/h6", ha="center", va="center", fontsize=9)
    ax.set_title("Mininet dumbbell topology used by the experiment runner", fontsize=13, weight="bold", pad=12)
    fig.tight_layout()
    fig.savefig(DOCS / "topology.png", dpi=240, bbox_inches="tight")
    plt.close(fig)


def set_cell_shading(cell, fill="DCE6F1"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=60, start=70, bottom=60, end=70):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def add_field(paragraph, field_code):
    run = paragraph.add_run()
    fldChar1 = OxmlElement("w:fldChar")
    fldChar1.set(qn("w:fldCharType"), "begin")
    instrText = OxmlElement("w:instrText")
    instrText.set(qn("xml:space"), "preserve")
    instrText.text = field_code
    fldChar2 = OxmlElement("w:fldChar")
    fldChar2.set(qn("w:fldCharType"), "end")
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def set_columns(section, count=2, space=360):
    sectPr = section._sectPr
    cols = sectPr.xpath("./w:cols")
    if cols:
        col = cols[0]
    else:
        col = OxmlElement("w:cols")
        sectPr.append(col)
    col.set(qn("w:num"), str(count))
    col.set(qn("w:space"), str(space))


def setup_document():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.58)
    sec.bottom_margin = Inches(0.58)
    sec.left_margin = Inches(0.62)
    sec.right_margin = Inches(0.62)
    sec.header_distance = Inches(0.25)
    sec.footer_distance = Inches(0.28)
    set_columns(sec, 1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(3)
    normal.paragraph_format.line_spacing = 1.02
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for style_name, size in (("Heading 1", 11), ("Heading 2", 9.8), ("Heading 3", 9.3)):
        st = styles[style_name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(30, 53, 75)
        st.paragraph_format.space_before = Pt(6 if style_name == "Heading 1" else 4)
        st.paragraph_format.space_after = Pt(2)
        st.paragraph_format.keep_with_next = True

    # Small running header and page-number footer.
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = header.add_run("TCP CONGESTION CONTROL | PRE-SUBMISSION DRAFT")
    r.font.name = "Times New Roman"; r.font.size = Pt(7); r.font.italic = True; r.font.color.rgb = RGBColor(100, 100, 100)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run("Page ")
    rr.font.name = "Times New Roman"; rr.font.size = Pt(8)
    add_field(footer, "PAGE")
    return doc


def add_text(doc, text, style=None, italic=False, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        run = p.add_run(text)
        run.italic = italic
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(8); r.italic = True
    return p


def add_table(doc, headers, rows, widths=None, font_size=7.3):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for i, text in enumerate(headers):
        cell = header_row.cells[i]
        cell.text = str(text)
        set_cell_shading(cell, "DCE6F1")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.font.name = "Times New Roman"; run.font.size = Pt(font_size); run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    run.font.name = "Times New Roman"; run.font.size = Pt(font_size)
    if widths:
        # Explicitly constrain both the table width and its grid. Setting only
        # cell.width leaves Word's default full-page grid, which clips in columns.
        total_twips = int(round(sum(widths) * 1440))
        tbl_pr = table._tbl.tblPr
        tbl_w = tbl_pr.find(qn("w:tblW"))
        if tbl_w is None:
            tbl_w = OxmlElement("w:tblW")
            tbl_pr.append(tbl_w)
        tbl_w.set(qn("w:w"), str(total_twips))
        tbl_w.set(qn("w:type"), "dxa")
        jc = tbl_pr.find(qn("w:jc"))
        if jc is None:
            jc = OxmlElement("w:jc")
            tbl_pr.append(jc)
        jc.set(qn("w:val"), "left")
        grid_cols = table._tbl.tblGrid.gridCol_lst
        for i, width in enumerate(widths):
            twips = int(round(width * 1440))
            if i < len(grid_cols):
                grid_cols[i].set(qn("w:w"), str(twips))
            for row in table.rows:
                row.cells[i].width = Inches(width)
                tc_pr = row.cells[i]._tc.get_or_add_tcPr()
                tc_w = tc_pr.find(qn("w:tcW"))
                if tc_w is None:
                    tc_w = OxmlElement("w:tcW")
                    tc_pr.append(tc_w)
                tc_w.set(qn("w:w"), str(twips))
                tc_w.set(qn("w:type"), "dxa")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_result_placeholder(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.06)
    p.paragraph_format.right_indent = Inches(0.06)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "FFF2CC"); pPr.append(shd)
    r = p.add_run(text)
    r.font.name = "Times New Roman"; r.font.size = Pt(8.2); r.font.bold = True; r.font.color.rgb = RGBColor(120, 70, 0)
    return p


def build():
    create_figures()
    doc = setup_document()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("An Experimental Comparative Analysis of TCP Reno, CUBIC, and BBR Under Varying Network Conditions")
    run.font.name = "Times New Roman"; run.font.size = Pt(17); run.bold = True; run.font.color.rgb = RGBColor(22, 48, 74)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Aditya Singh Rana, Vinay Mor, Tejas Pisal")
    r.font.name = "Times New Roman"; r.font.size = Pt(10); r.bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("[Confirm department, institution, and author contact details before submission]")
    r.font.name = "Times New Roman"; r.font.size = Pt(8); r.italic = True; r.font.color.rgb = RGBColor(90, 90, 90)

    add_result_placeholder(doc, "AUTHOR ACTION REQUIRED — This is a pre-submission research draft. No experimental measurements have been run in the current authoring environment. Run the supplied experiment package, replace every bracketed results item, update the conclusion, remove this note, and apply the target conference template before submission. Do not present placeholder values as results.")

    add_heading(doc, "Abstract", 1)
    add_text(doc,
        "TCP congestion-control algorithms can produce different throughput, delay, retransmission, and bandwidth-sharing behavior under the same network conditions. This paper specifies a reproducible experimental comparison of Linux TCP Reno, CUBIC, and BBR using a Mininet dumbbell topology with a shared bottleneck. The supplied experiment artifact automates one-factor-at-a-time variations in bottleneck bandwidth, configured link delay, random packet loss, and queue limit, and includes a mixed-flow experiment in which Reno, CUBIC, and BBR compete concurrently. Receiver-side throughput is collected using iperf3; sender retransmissions are extracted from iperf3 summaries; and TCP RTT and congestion-window samples are collected from Linux socket statistics where available. Raw JSON, per-run CSV data, environment metadata, confidence-interval summaries, and plots are retained to support auditability. Because the experiments have not yet been executed in the target Linux environment, this draft does not claim numerical performance findings. The final abstract must be updated from measured results and must state the principal quantitative findings rather than anticipated outcomes.")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Index Terms—")
    r.bold = True; r.font.name = "Times New Roman"; r.font.size = Pt(9)
    r2 = p.add_run("TCP congestion control, Reno, CUBIC, BBR, Mininet, throughput, round-trip time, fairness, reproducibility.")
    r2.font.name = "Times New Roman"; r2.font.size = Pt(9)

    # Continuous section with two columns for the main paper.
    sec2 = doc.add_section(WD_SECTION_START.CONTINUOUS)
    sec2.page_width = Inches(8.27); sec2.page_height = Inches(11.69)
    sec2.top_margin = Inches(0.58); sec2.bottom_margin = Inches(0.58)
    sec2.left_margin = Inches(0.62); sec2.right_margin = Inches(0.62)
    sec2.header_distance = Inches(0.25); sec2.footer_distance = Inches(0.28)
    set_columns(sec2, 2, 360)
    # Keep header/footer matching after section split.
    sec2.header.is_linked_to_previous = True
    sec2.footer.is_linked_to_previous = True

    add_heading(doc, "I. Introduction", 1)
    add_text(doc,
        "Transmission Control Protocol (TCP) provides reliable, ordered delivery of byte streams, while congestion control limits the rate at which a sender injects data into the network [1], [2]. The congestion-control algorithm influences how quickly a connection uses available capacity, responds to congestion signals, recovers after loss, and shares a bottleneck with other flows. These effects depend on more than the algorithm label: kernel implementation, path delay, buffer configuration, traffic mix, and competing connections can change the observed outcome.")
    add_text(doc,
        "Reno represents the classical loss-based additive-increase/multiplicative-decrease family described by TCP congestion-control standards [2]. CUBIC changes congestion-window growth using a cubic function and is specified in RFC 9438 [3]. BBR instead estimates bottleneck bandwidth and minimum round-trip propagation time to pace traffic using a model of the path [4]. These design differences motivate a comparison that measures multiple metrics and avoids declaring one algorithm universally superior.")
    add_text(doc,
        "Prior evaluations have already examined BBR throughput, queuing delay, loss, and fairness [5], as well as the fairness of BBR and CUBIC flows sharing a bottleneck [6]. BBRv3 has also been evaluated under variations in buffer size, flow count, delay, and loss [7]. More recently, cloud-based work compared Reno, CUBIC, and BBR across synthetic throughput tests and application workloads [8]. Therefore, this study does not claim that a three-algorithm comparison is itself novel. Its intended contribution is a compact, reproducible laboratory artifact that applies a documented one-factor-at-a-time matrix, includes a mixed Reno/CUBIC/BBR sharing case, preserves raw observations, and generates summaries from the same data used in the manuscript.")
    add_text(doc,
        "The research questions are: (RQ1) How do receiver-side throughput, sampled TCP RTT, retransmissions, and congestion-window observations change as bandwidth, configured delay, loss, and queue limits vary? (RQ2) How is throughput divided when one Reno, one CUBIC, and one BBR flow concurrently share the same bottleneck? (RQ3) Can the resulting measurements, run settings, and analysis outputs be generated from a documented and repeatable procedure?")

    add_heading(doc, "II. Background and Related Work", 1)
    add_heading(doc, "A. Congestion-control mechanisms", 2)
    add_text(doc,
        "Reno uses loss indications to trigger window reduction and then increases its congestion window conservatively during congestion avoidance [2]. CUBIC defines a cubic window-growth function based on elapsed time since the last congestion event; the modern standards-track specification is RFC 9438, which updates RFC 5681 in its treatment of CUBIC behavior [3]. BBR estimates bottleneck delivery rate and minimum RTT and uses those estimates to control pacing rather than relying primarily on packet loss as its congestion signal [4]. The Linux system under test determines the concrete implementation of each algorithm; results must therefore report the kernel release and the congestion-control names exposed by the kernel.")
    add_heading(doc, "B. Prior experimental evaluations", 2)
    add_text(doc,
        "Hock, Bless, and Zitterbart independently evaluated BBR with high-rate bottlenecks, varying RTT, flow count, and buffer size. Their study reported the intended utilization behavior but also observed cases involving increased queuing delay, unfairness, and packet loss [5]. Sasaki et al. investigated concurrent BBR and CUBIC connections and reported substantial fairness problems in some high-latency conditions [6]. These findings motivate measuring both per-flow outcomes and aggregate performance, rather than throughput alone.")
    add_text(doc,
        "Gomez et al. evaluated BBRv3 alongside CUBIC under different flow counts, propagation delays, loss rates, buffer sizes, and queue-management schemes [7]. Swamy et al. subsequently compared Reno, CUBIC, and BBR in an AWS dumbbell environment across synthetic throughput and application workloads, illustrating that congestion-control behavior can differ across workload types [8]. Mininet enables small-scale network emulation using lightweight OS-level virtualization [9], but an emulated result remains conditional on the kernel, host scheduling, virtualization, and configured topology. The present work is best understood as a reproducibility-oriented controlled study, not as evidence that existing large-scale or workload-diverse studies are absent.")

    add_heading(doc, "III. Study Design and Contributions", 1)
    add_text(doc,
        "The study compares the Linux implementations identified by the kernel as `reno`, `cubic`, and `bbr`. It examines receiver-side throughput as the primary performance measure and uses sender retransmission counts, sampled TCP RTT, congestion-window samples, and mixed-flow fairness as complementary measures. The artifact contributes (i) an executable Mininet testbed with a single shared bottleneck, (ii) a documented scenario matrix that varies one primary impairment at a time, (iii) persistent raw and processed outputs, and (iv) an analysis script that reports mean, standard deviation, sample count, and 95% confidence intervals for repeated observations.")
    add_text(doc,
        "The contribution is methodological and empirical. A claim of algorithmic novelty is not made. Whether the resulting study is sufficiently novel for a particular conference must be assessed against that venue's scope and the latest literature, especially because Reno/CUBIC/BBR comparisons and fairness studies already exist [5]–[8].")

    add_heading(doc, "IV. Experimental Methodology", 1)
    add_heading(doc, "A. Testbed and topology", 2)
    add_text(doc,
        "The provided implementation creates six Mininet hosts and two Open vSwitch switches. Three senders are attached to the left switch, three receivers are attached to the right switch, and all sender-to-receiver traffic crosses one configurable inter-switch bottleneck. Access links are configured above the bottleneck rate. The topology supports both isolated single-flow trials and three simultaneous competing flows. The design follows the common dumbbell arrangement used to make a shared contention point observable [5], [6], [8], [9].")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(DOCS / "architecture.png"), width=Inches(3.18))
    add_caption(doc, "Fig. 1. Experiment automation, measurement, data retention, and analysis pipeline.")
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(DOCS / "topology.png"), width=Inches(3.18))
    add_caption(doc, "Fig. 2. Mininet dumbbell topology. The bottleneck is shared by all flows.")
    add_text(doc,
        "The script writes the kernel release, tool version, selected congestion-control algorithm names, scenario parameters, and experiment timestamps to the output directory. Before collecting publishable measurements, the operator must inspect the actual qdisc settings (for example, with `tc -s qdisc show` inside the relevant namespaces) and retain the environment file with the raw logs. Delay is controlled using Mininet's TCLink delay parameter; the configured parameter must not be reported as an observed end-to-end RTT. RTT is measured separately from sampled socket statistics.")

    add_heading(doc, "B. Experimental parameters", 2)
    add_text(doc,
        "A one-factor-at-a-time matrix varies bottleneck bandwidth, link delay, injected random loss, and queue limit while holding the other listed settings at their baseline values. The numerical settings below are the proposed default values in the supplied runner; they are experimental design choices, not measured findings. A full execution uses 30-second transfers, omits the initial 3 seconds from iperf3's measurement interval, and repeats each algorithm/scenario combination five times. The script randomizes algorithm order within each scenario/repetition using a recorded seed.")
    add_table(doc,
        ["Scenario", "Varied setting", "Settings held at baseline"],
        [
            ["Baseline", "None", "20 Mbit/s; 10 ms; 0% loss; 100-packet queue"],
            ["Low bandwidth", "5 Mbit/s", "10 ms; 0% loss; 100-packet queue"],
            ["High delay", "50 ms", "20 Mbit/s; 0% loss; 100-packet queue"],
            ["Random loss", "0.5% loss", "20 Mbit/s; 10 ms; 100-packet queue"],
            ["Small queue", "20 packets", "20 Mbit/s; 10 ms; 0% loss"],
        ], widths=[0.85, 0.8, 1.65], font_size=6.7)
    add_caption(doc, "TABLE I. Default one-factor-at-a-time experimental matrix configured by the supplied artifact.")
    add_text(doc,
        "For single-flow experiments, h1 sends a bulk TCP transfer to h4. The sender is configured with one of the three congestion-control algorithms before each trial. For mixed-flow fairness, h1, h2, and h3 simultaneously transfer to h4, h5, and h6, respectively, using Reno, CUBIC, and BBR. A mixed-flow repetition shares a common run identifier so that throughputs are grouped correctly for fairness analysis. The test runner records one output row per flow; raw iperf3 JSON is retained independently.")

    add_heading(doc, "C. Metrics and measurement procedure", 2)
    add_text(doc,
        "Receiver-side throughput is calculated from iperf3's `end.sum_received.bits_per_second` field, converted to Mbit/s. Sender throughput is retained as a cross-check. Sender retransmissions are extracted from the iperf3 sender summary when exposed by the installed version. These counts are not equivalent to the configured random-loss percentage: retransmissions depend on transport behavior and measurement details.")
    add_text(doc,
        "TCP RTT and congestion-window values are sampled from `ss -tin` while a flow is active. The runner stores the mean, median, and 95th percentile of captured RTT samples, and the median and maximum captured congestion-window values. The sampling interval is approximately 350 ms; fields may be absent on systems whose socket reporting differs or if the connection is too short to sample. Missing values remain empty rather than being replaced with zero. The standard iperf3 jitter field is not used as a TCP jitter measurement; a separate UDP workload would be needed if jitter were a research objective.")
    add_text(doc,
        "For a mixed-flow run with throughputs x1, ..., xn, Jain's fairness index is computed as J = (sum(xi))^2 / [n sum(xi^2)] [10]. The index lies between 1/n and 1 for positive allocations, with 1 indicating equal throughput across the measured flows. Fairness is interpreted alongside total throughput because equal shares do not, by themselves, imply high link utilization.")

    add_heading(doc, "D. Repetition and analysis", 2)
    add_text(doc,
        "The experiment runner uses repeated independent runs for each condition. The analysis program aggregates each scenario/algorithm group and reports the number of valid observations, mean, sample standard deviation, and a two-sided 95% t-based confidence interval for metrics with sufficient non-missing values. For mixed-flow tests, the analysis computes one fairness index per simultaneous-run identifier. The generated figures include mean values with confidence intervals where calculable. Any exclusion of failed or incomplete runs must be documented; the default analysis does not silently replace missing measurements.")
    add_text(doc,
        "Run ordering, host workload, kernel version, queue discipline, CPU scheduling, segmentation/offload settings, and virtualization configuration can influence the result. The code records the kernel and tool versions but does not fully normalize every hardware or offload property. Before the final study, the authors should record `ethtool -k` output where relevant, inspect qdiscs, keep the machine otherwise idle, and report any configuration changes. The analysis is descriptive and does not automatically establish statistical significance between algorithms.")

    add_heading(doc, "V. Results and Discussion", 1)
    add_result_placeholder(doc,
        "RESULTS REQUIRED — This section must be completed after running `src/run_experiments.py` on the documented Linux/Mininet environment. Use `src/analyze_results.py` to create the measured summaries and plots. Do not submit this manuscript while the bracketed items and this instruction remain.")
    add_heading(doc, "A. Single-flow performance", 2)
    add_text(doc,
        "Report receiver-side throughput by algorithm and scenario using the generated `summary_by_scenario_algorithm.csv` and the corresponding throughput plot. Include observed sample counts and 95% confidence intervals. Then explain the measured changes across the bandwidth, delay, loss, and queue settings without generalizing beyond the test conditions.")
    add_result_placeholder(doc, "[INSERT: measured throughput table and Figure 3 generated from the full experimental runs. Summarize actual differences with units, n, and uncertainty.]")
    add_heading(doc, "B. Delay, retransmissions, and congestion window", 2)
    add_text(doc,
        "Use the RTT, retransmission, and congestion-window plots only where the raw logs contain valid measurements. Compare trends with the algorithm designs and prior evaluations [4]–[7], while noting that `ss` sampling is periodic rather than packet-by-packet tracing. Explain any absent or anomalous measurements and verify them against the corresponding raw JSON and system logs before drawing conclusions.")
    add_result_placeholder(doc, "[INSERT: measured RTT/retransmission/cwnd findings and figures. State how many observations were available for each metric; do not treat configured loss as measured packet loss.]")
    add_heading(doc, "C. Mixed-flow fairness", 2)
    add_text(doc,
        "For the mixed-flow experiment, report each flow's measured throughput, aggregate throughput, and Jain's fairness index across repeated concurrent runs. Explain which flow received what share under this particular bottleneck, and compare the result cautiously with prior BBR/CUBIC fairness studies [5]–[7].")
    add_result_placeholder(doc, "[INSERT: measured mixed-flow fairness table and `jain_fairness.png`; report n and confidence interval.]")
    add_heading(doc, "D. Interpretation", 2)
    add_text(doc,
        "The final discussion should answer RQ1–RQ3 directly. Distinguish observations from explanations: measured throughput or RTT is an observation, while attributing it to pacing, loss response, or window growth is an interpretation that should be supported by the recorded configuration and relevant literature. If outcomes vary across repetitions or conflict with expected behavior, report the variation rather than discarding inconvenient runs without a documented rule.")

    add_heading(doc, "VI. Threats to Validity and Limitations", 1)
    add_text(doc,
        "First, Mininet shares a host kernel and physical CPU resources among virtual nodes; its timing and throughput may be affected by host scheduling and virtualization overhead. Second, results depend on the specific Linux kernel's Reno, CUBIC, and BBR implementations. The label `bbr` alone does not establish that the tested algorithm is BBRv1, BBRv2, or BBRv3; the version should be verified from kernel documentation or build metadata where possible. Third, the matrix uses a limited set of one-factor-at-a-time configurations and does not cover every interaction between bandwidth, delay, loss, buffer size, active queue management, and flow count. Fourth, the short set of bulk iperf3 transfers is not representative of all application workloads, wireless paths, Internet routes, or short-flow behavior. Finally, socket RTT/cwnd sampling can miss fast transient events. Accordingly, conclusions should be restricted to the configured emulated testbed and verified measured results.")

    add_heading(doc, "VII. Conclusion", 1)
    add_result_placeholder(doc,
        "CONCLUSION REQUIRED — Replace this instruction after the experiment run. State the key measured finding for each research question, quantify the most important differences, acknowledge uncertainty and scope, and identify one or two evidence-based future-work directions. Do not state that BBR, CUBIC, or Reno is universally superior.")

    add_heading(doc, "Artifact and Data Availability", 1)
    add_text(doc,
        "The experiment source code, configuration matrix, analysis script, and instructions are supplied with this manuscript draft. For a public conference submission, archive the code and final raw/processed dataset in a stable repository and replace this sentence with the permanent URL and version/commit identifier. Do not publish sensitive host data; the intended artifacts here are network metrics, tool/kernel metadata, and experiment outputs.")

    add_heading(doc, "References", 1)
    refs = [
        "[1] W. Eddy, Ed., “Transmission Control Protocol (TCP),” RFC 9293, Aug. 2022, doi: 10.17487/RFC9293.",
        "[2] M. Allman, V. Paxson, and E. Blanton, “TCP Congestion Control,” RFC 5681, Sep. 2009, updated by RFC 9438.",
        "[3] L. Xu, S. Ha, I. Rhee, V. Goel, and L. Eggert, “CUBIC for Fast and Long-Distance Networks,” RFC 9438, Aug. 2023.",
        "[4] N. Cardwell, Y. Cheng, C. S. Gunn, S. H. Yeganeh, and V. Jacobson, “BBR: Congestion-Based Congestion Control,” Communications of the ACM, vol. 60, no. 2, pp. 58–66, 2017, doi: 10.1145/3009824.",
        "[5] M. Hock, R. Bless, and M. Zitterbart, “Experimental evaluation of BBR congestion control,” in Proc. IEEE 25th International Conference on Network Protocols (ICNP), 2017, pp. 1–10, doi: 10.1109/ICNP.2017.8117540.",
        "[6] K. Sasaki, M. Hanai, K. Miyazawa, A. Kobayashi, N. Oda, and S. Yamaguchi, “TCP Fairness Among Modern TCP Congestion Control Algorithms Including TCP BBR,” in Proc. IEEE 7th International Conference on Cloud Networking (CloudNet), 2018, pp. 1–4, doi: 10.1109/CloudNet.2018.8549505.",
        "[7] J. Gomez, E. F. Kfoury, J. Crichigno, and G. Srivastava, “Evaluating TCP BBRv3 performance in wired broadband networks,” Computer Communications, vol. 222, pp. 198–208, Jun. 2024, doi: 10.1016/j.comcom.2024.04.037.",
        "[8] Aravinda Swamy D. et al., “Performance evaluation of TCP congestion control variants across application workloads in cloud based networks,” Scientific Reports, vol. 16, Art. no. 31639, Jul. 2026, doi: 10.1038/s41598-026-61100-w.",
        "[9] B. Lantz, B. Heller, and N. McKeown, “A network in a laptop: rapid prototyping for software-defined networks,” in Proc. 10th ACM Workshop on Hot Topics in Networks (HotNets), 2010, doi: 10.1145/1868447.1868466.",
        "[10] R. Jain, D. Chiu, and W. Hawe, “A Quantitative Measure of Fairness and Discrimination for Resource Allocation in Shared Computer Systems,” DEC Research Report TR-301, Sep. 1984.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.first_line_indent = Inches(-0.12)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(ref)
        run.font.name = "Times New Roman"; run.font.size = Pt(7.7)

    output = ROOT / "TCP_Congestion_Control_Conference_Draft_PENDING_RESULTS.docx"
    doc.save(output)
    print(output)

if __name__ == "__main__":
    build()
